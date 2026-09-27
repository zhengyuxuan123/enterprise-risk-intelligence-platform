"""分析报告导出（Word / PDF）—— 对应 Java ``ReportExportService``。

全部本地渲染，不依赖任何外部服务。

**一条硬要求：导出必须带上来源。**
报告一旦离开系统就会在邮件和群里流转，如果导出时把引用编号和来源 URL 丢了，
它就变成一份"看起来很专业但无从核对"的结论 —— 这与整套系统的溯源主张直接冲突。
所以这里强制输出「参考来源（网页）」清单（标题 + 完整 URL）与知识库来源标题，
并附上引用核对摘要，让拿到报告的人能自行判断可信度。

Java 用 Apache POI + PDFBox，这里换成 ``python-docx`` + ``reportlab``。
排版参数（字号、行距、页边距）与 Java 保持一致，让两份产物看上去是同一份报告。
"""

from __future__ import annotations

import io
import json
import logging
import os
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from sqlalchemy import select

from ..core.errors import BusinessError
from ..db import get_db
from ..db import tables as T

log = logging.getLogger(__name__)

DTF = "%Y-%m-%d %H:%M"

#: PDF 中文字体的候选（与 Java 同序）。PDF 不嵌入字体就无法正确显示中文。
_FONT_DIRS = [os.environ.get("SystemRoot", "") + "\\Fonts",
              "C:\\Windows\\Fonts", "/usr/share/fonts/truetype", "/usr/share/fonts"]
_FONT_NAMES = ["simhei.ttf", "SimHei.ttf", "simkai.ttf", "SIMFANG.TTF", "simfang.ttf",
               "msyh.ttf", "Deng.ttf", "NotoSansCJK-Regular.ttc"]


@dataclass
class Report:
    filename: str
    content_type: str
    bytes: bytes


@dataclass
class Line:
    style: str
    text: str


def _str(o: Any) -> str:
    return "" if o is None else str(o).strip()


def _null_to_empty(s: Optional[str]) -> str:
    return (s or "").strip()


def _parse(js: Any) -> Dict[str, Any]:
    if not js:
        return {}
    if isinstance(js, dict):
        return js
    try:
        m = json.loads(js)
        return m if isinstance(m, dict) else {}
    except Exception:  # noqa: BLE001
        return {}


def _list_of(o: Any) -> List[Dict[str, Any]]:
    if not isinstance(o, list):
        return []
    return [x for x in o if isinstance(x, dict)]


def _recommendations(s: Dict[str, Any], key: str) -> List[str]:
    """建议可能是 ``List[{action:...}]`` 也可能是 ``List[str]``，两种都认。"""
    o = s.get(key)
    if not isinstance(o, list):
        return []
    out: List[str] = []
    for x in o:
        if isinstance(x, dict):
            v = x.get("action")
            if _str(v):
                out.append(_str(v))
        elif x is not None and _str(x):
            out.append(_str(x))
    return out


class ReportExportService:
    def __init__(self, db=None, scope=None) -> None:
        self.db = db
        #: 导出同样要过数据权限：不能凭一个 ID 就把别家企业的报告拿走。
        self.scope = scope

    # ------------------------------------------------------------------

    def export(self, analysis_id: int, fmt: Optional[str] = None) -> Report:
        db = self._db()
        a = db.fetch_one(select(T.ai_analysis).where(T.ai_analysis.c.id == analysis_id))
        if a is None:
            raise BusinessError(f"分析记录不存在：{analysis_id}")
        if self.scope is not None and a.get("company_id") is not None \
                and not self.scope(a["company_id"]):
            raise BusinessError("不能导出其他企业的分析报告")

        company_name = "未知企业"
        try:
            c = db.fetch_one(select(T.company).where(T.company.c.id == a.get("company_id")))
            if c and c.get("company_name"):
                company_name = c["company_name"]
        except Exception:  # noqa: BLE001
            pass

        lines = self.build_lines(a, company_name)
        is_docx = (fmt or "").lower() == "docx"
        base = f"风险分析报告-{company_name}-{a.get('id')}"
        try:
            if is_docx:
                return Report(base + ".docx",
                              "application/vnd.openxmlformats-officedocument"
                              ".wordprocessingml.document",
                              self.to_docx(lines))
            return Report(base + ".pdf", "application/pdf", self.to_pdf(lines))
        except BusinessError:
            raise
        except Exception as e:  # noqa: BLE001
            # 最常见的失败是找不到中文字体，说清楚原因，别让用户对着 500 猜
            raise BusinessError(f"导出失败：{e}") from e

    # ------------------------------------------------------------------
    # 内容组装
    # ------------------------------------------------------------------

    def build_lines(self, a: Dict[str, Any], company_name: str) -> List[Line]:
        out: List[Line] = []
        s = _parse(a.get("answer_json"))

        out.append(Line("h1", "企业经营风险分析报告"))
        out.append(Line("meta", "企业：" + company_name))
        out.append(Line("meta", "分析编号：#" + _str(a.get("id"))))
        out.append(Line("meta", "提问：" + _null_to_empty(a.get("question"))))
        created = a.get("created_at")
        out.append(Line("meta", "生成时间：" + (created.strftime(DTF) if isinstance(created, datetime)
                                            else "—")))
        out.append(Line("meta", "风险等级：" + _str(s.get("risk_level"))))
        out.append(Line("meta", "置信度：" + _null_to_empty(a.get("confidence"))
                        + "　溯源：" + ("通过" if a.get("grounded") == 1 else "待人工核对")))
        if s.get("degraded") is True:
            out.append(Line("warn", "注意：本次为降级报告，模型服务不可用，结论基于本地数据规则生成。"))
        if s.get("review_revised") is True:
            out.append(Line("meta", "本报告已根据复核意见修订（触发原因："
                            + _str(s.get("review_trigger")) + "）"))
        out.append(Line("gap", ""))

        headline = _str(s.get("headline"))
        if headline:
            out.append(Line("h2", "一、核心结论"))
            out.append(Line("p", headline))
            out.append(Line("gap", ""))

        out.append(Line("h2", "二、分析结论"))
        ci, cw = _str(s.get("conclusion_internal")), _str(s.get("conclusion_web"))
        if ci:
            out.append(Line("h3", "内部依据"))
            out.append(Line("p", ci))
        if cw:
            out.append(Line("h3", "外部参考"))
            out.append(Line("p", cw))
        out.append(Line("gap", ""))

        out.append(Line("h2", "三、处置建议"))
        n = 0
        for r in _recommendations(s, "recommendations_internal"):
            out.append(Line("bullet", "（内部）" + r))
            n += 1
        for r in _recommendations(s, "recommendations_web"):
            out.append(Line("bullet", "（外部）" + r))
            n += 1
        if n == 0:
            out.append(Line("p", "本次未给出具体动作建议。"))
        out.append(Line("gap", ""))

        # ---- 来源：导出的底线，不能省 ----
        web = _list_of(s.get("web_sources"))
        kb = _list_of(s.get("kb_sources"))
        out.append(Line("h2", "四、参考来源"))
        if web:
            out.append(Line("h3", f"网页来源（{len(web)} 条）"))
            for i, w in enumerate(web, 1):
                url = _str(w.get("url"))
                out.append(Line("bullet", f"[{i}] " + _str(w.get("title"))))
                if url:
                    out.append(Line("small", "    " + url))
        if kb:
            out.append(Line("h3", f"知识库来源（{len(kb)} 条）"))
            for k in kb:
                did = _str(k.get("sourceRef")) or _str(k.get("documentId"))
                prefix = (f"来源ID={did}　" if did else "")
                out.append(Line("bullet", prefix + _str(k.get("title"))))
        if not web and not kb:
            out.append(Line("p", "本次分析未引用可列出的来源。"))
        out.append(Line("gap", ""))

        gs = _str(s.get("grounding_summary"))
        if gs:
            out.append(Line("h2", "五、引用核对"))
            out.append(Line("p", gs))
            out.append(Line("gap", ""))

        answer = a.get("answer")
        if answer and str(answer).strip():
            out.append(Line("h2", "附录：完整正文"))
            for para in str(answer).split("\n"):
                t = para.strip()
                if t:
                    out.append(Line("p", t))
        return out

    # ------------------------------------------------------------------
    # Word
    # ------------------------------------------------------------------

    def to_docx(self, lines: List[Line]) -> bytes:
        from docx import Document
        from docx.shared import Pt, RGBColor

        doc = Document()
        for l in lines:
            if l.style == "gap":
                continue
            p = doc.add_paragraph()
            run = p.add_run(("· " + l.text) if l.style == "bullet" else l.text)
            run.font.name = "宋体"
            if l.style == "h1":
                run.bold, run.font.size = True, Pt(18)
            elif l.style == "h2":
                run.bold, run.font.size = True, Pt(14)
            elif l.style == "h3":
                run.bold, run.font.size = True, Pt(12)
            elif l.style == "warn":
                run.bold, run.font.size = True, Pt(11)
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            elif l.style in ("meta", "small"):
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            else:
                run.font.size = Pt(11)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------

    def to_pdf(self, lines: List[Line]) -> bytes:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        font_name, font_path = self._load_chinese_font()
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        page_w, page_h = A4
        margin = 50.0
        max_w = page_w - margin * 2
        y = page_h - margin

        for l in lines:
            if l.style == "gap":
                y -= 10
                continue
            size = {"h1": 18.0, "h2": 14.0, "h3": 12.0,
                    "meta": 9.0, "small": 9.0}.get(l.style, 11.0)
            leading = size * 1.6
            color = {"h1": (0.08, 0.08, 0.08), "h2": (0.08, 0.08, 0.08),
                     "warn": (0.71, 0.12, 0.12),
                     "meta": (0.43, 0.43, 0.43), "small": (0.43, 0.43, 0.43)}.get(
                         l.style, (0.12, 0.12, 0.12))
            for seg in self._wrap(l.text, font_name, size, max_w):
                if y < margin + leading:
                    c.showPage()
                    y = page_h - margin
                y -= leading
                c.setFont(font_name, size)
                c.setFillColorRGB(*color)
                c.drawString(margin, y, seg)
        c.showPage()
        c.save()
        return buf.getvalue()

    @staticmethod
    def _wrap(text: str, font_name: str, size: float, max_w: float) -> List[str]:
        """按字符宽度贪心换行；中英文混排下逐字符累加最稳。"""
        from reportlab.pdfbase.pdfmetrics import stringWidth

        if not text:
            return []
        out: List[str] = []
        cur = ""
        for ch in text:
            if ch == "\n":
                out.append(cur)
                cur = ""
                continue
            if cur and stringWidth(cur + ch, font_name, size) > max_w:
                out.append(cur)
                cur = ch
            else:
                cur += ch
        if cur:
            out.append(cur)
        return out

    @staticmethod
    def _load_chinese_font() -> Tuple[str, str]:
        """找一个系统里存在的中文字体。找不到就把原因说清楚。"""
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        for d in _FONT_DIRS:
            if not d:
                continue
            for n in _FONT_NAMES:
                p = os.path.join(d, n)
                if not os.path.exists(p):
                    continue
                try:
                    pdfmetrics.registerFont(TTFont(n, p))
                    return n, p
                except Exception:  # noqa: BLE001
                    continue
        raise BusinessError(
            "未找到可用的中文字体（已尝试 " + "/".join(_FONT_NAMES)
            + "）。请在服务器上安装中文字体后重试，或改用 Word 导出。")

    def _db(self):
        if self.db is None:
            self.db = get_db()
        return self.db
